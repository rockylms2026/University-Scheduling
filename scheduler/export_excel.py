"""Excel workbook: timetables by student group, teacher, and room, plus lists."""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .model import University, fmt_time
from .solver import Schedule, Session
from .store import OUTPUT_DIR
from .validate import faculty_load_summary, validate
from .verify import verify

COLLEGE_FILL = {
    "CSS": "DCE9F7", "CISC": "DDF0E1", "CBE": "FCEBD3", "COE": "EADFF5",
    "CCJPS": "F8DCDC", "CORE": "EEEEEE",
}
HEAD = PatternFill("solid", fgColor="2F3E4E")
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _days(uni: University) -> list[str]:
    return uni.timeframe.days


def _slot_times(uni: University) -> list[int]:
    tf = uni.timeframe
    return list(range(tf.day_start, tf.day_end, tf.slot))


def _cell_text(s: Session, view: str) -> str:
    k = s.klass
    lines = [f"{k.course.code} {s.part if s.part not in ('Lec',) else ''}".strip()]
    if view != "section":
        lines.append(k.section.name)
    if view != "room" and s.room:
        lines.append(s.room)
    if view != "teacher":
        lines.append(k.faculty.short if k.faculty else "TBA")
    if s.exception:
        lines.append("(needs approval)")
    return "\n".join(lines)


def _grid(ws, row: int, title: str, subtitle: str, sessions: list[Session], uni: University, view: str) -> int:
    """Write one weekly grid starting at row; return the next free row."""
    days, times = _days(uni), _slot_times(uni)
    ncol = len(days) + 1
    ws.cell(row, 1, title).font = Font(bold=True, size=12)
    ws.cell(row + 1, 1, subtitle).font = Font(italic=True, color="555555")
    head = row + 2
    ws.cell(head, 1, "Time").font = Font(bold=True, color="FFFFFF")
    ws.cell(head, 1).fill = HEAD
    for j, d in enumerate(days, start=2):
        c = ws.cell(head, j, d)
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), HEAD, CENTER
    for i, t in enumerate(times):
        c = ws.cell(head + 1 + i, 1, f"{fmt_time(t)}-{fmt_time(t + uni.timeframe.slot)}")
        c.font, c.border = Font(size=8, color="555555"), BOX
        for j in range(2, ncol + 1):
            ws.cell(head + 1 + i, j).border = BOX
    taken: dict[tuple[int, int], int] = {}      # (row, col) -> anchor row of its block
    for s in sorted(sessions, key=lambda x: (x.day, x.start)):
        if s.day not in days:
            continue
        col = days.index(s.day) + 2
        r0 = head + 1 + (s.start - uni.timeframe.day_start) // uni.timeframe.slot
        r1 = head + (s.end - uni.timeframe.day_start) // uni.timeframe.slot
        overlap = next((taken[(r, col)] for r in range(r0, r1 + 1) if (r, col) in taken), None)
        if overlap is not None:
            # A venue holding two classes at once: list both in one block.
            anchor = ws.cell(overlap, col)
            anchor.value = f"{anchor.value}\n---\n{_cell_text(s, view)}"
            continue
        cell = ws.cell(r0, col, _cell_text(s, view))
        cell.alignment = CENTER
        cell.font = Font(size=8, bold=bool(s.exception))
        color = "FFF2A8" if s.exception else COLLEGE_FILL.get(s.klass.course.college, "EEEEEE")
        for r in range(r0, r1 + 1):
            ws.cell(r, col).fill = PatternFill("solid", fgColor=color)
            taken[(r, col)] = r0
        if r1 > r0:
            ws.merge_cells(start_row=r0, start_column=col, end_row=r1, end_column=col)
    return head + len(times) + 3


def _setup_grid_sheet(ws, uni: University):
    ws.column_dimensions["A"].width = 13
    for j in range(2, len(_days(uni)) + 2):
        ws.column_dimensions[get_column_letter(j)].width = 19


def _table(ws, headers: list[str], rows: list[list], widths: list[int] | None = None):
    for j, h in enumerate(headers, start=1):
        c = ws.cell(1, j, h)
        c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), HEAD, CENTER
    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            ws.cell(i, j, v).alignment = WRAP
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for j, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w


def _meeting_rows(uni: University, sched: Schedule) -> list[list]:
    grouped: dict[str, list[Session]] = defaultdict(list)
    for s in sched.sessions:
        grouped[s.meeting_id].append(s)
    rows = []
    order = {d: i for i, d in enumerate(_days(uni))}
    for mid, ss in grouped.items():
        ss.sort(key=lambda s: order[s.day])
        s = ss[0]
        k = s.klass
        days = "".join("Th" if x.day == "Thu" else x.day[0] if x.day != "Sat" else "Sat" for x in ss)
        rooms = sorted({x.room or "-" for x in ss})
        rows.append([k.course.code, k.course.title, k.section.name, k.section.college, s.part, days,
                     f"{fmt_time(s.start)}-{fmt_time(s.end)}", s.label, ", ".join(rooms),
                     k.teacher_label, k.faculty.college if k.faculty else "",
                     "GE" if k.course.is_core else "Major", s.exception or ""])
    rows.sort(key=lambda r: (r[3], r[2], r[0], r[4]))
    return rows


def write_workbook(uni: University, schedules: dict[int, Schedule], path: Path | None = None) -> Path:
    path = path or OUTPUT_DIR / "University_Schedule.xlsx"
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    lines = [
        ("University Class Schedule", True),
        ("Built from the college curricula and faculty loading plans, the University Class Time Frame, "
         "and the Campus Space Plan.", False),
        ("", False),
        ("Sheets", True),
        ("Sem N Groups: weekly timetable of every student group (block section, or block and major/track).", False),
        ("Sem N Teachers: weekly timetable of every faculty member.", False),
        ("Sem N Rooms: weekly use of every classroom, laboratory, and PE venue.", False),
        ("Sem N List: one row per class meeting; use the filters to sort by college, section, room, or teacher.", False),
        ("Loads: teaching load, preparations, and days on campus per faculty member.", False),
        ("Room use: share of each room's 39 weekly teaching hours that is booked.", False),
        ("Needs approval: classes placed outside the time frame because no compliant time exists.", False),
        ("Plan notes: points in the loading plans that need a decision.", False),
        ("", False),
        ("Colors: blue CSS, green CISC, orange CBE, purple COE, red CCJPS, gray GE; yellow = needs approval.", False),
        ("TBA: GE sections that no college or GE Department plan assigns to a named teacher.", False),
        ("", False),
        ("Checks", True),
    ]
    for sem, sched in schedules.items():
        problems = verify(uni, sched)
        q = sched.stats.get("quality", {})
        lines.append((f"Semester {sem}: {len(sched.sessions)} sessions; "
                      f"{'no conflicts found' if not problems else f'{len(problems)} problems'} "
                      f"(student groups, teachers, rooms, lunch, placement days, daily limits re-checked). "
                      f"Average idle time per student group: {q.get('avg_idle_hours_per_student_group_per_week', '?')} h/week.",
                      False))
    for i, (text, bold) in enumerate(lines, start=1):
        readme.cell(i, 1, text).font = Font(bold=bold, size=14 if i == 1 else 11)
    readme.column_dimensions["A"].width = 130

    for sem, sched in schedules.items():
        # Student groups (subcohorts) include their block and track classes.
        ws = wb.create_sheet(f"Sem {sem} Groups")
        _setup_grid_sheet(ws, uni)
        by_sub = defaultdict(list)
        for s in sched.sessions:
            if s.part != "Supervision":
                for sub in s.klass.section.subcohorts:
                    by_sub[sub].append(s)
        row = 1
        for sub in sorted(by_sub, key=lambda x: (uni.programs[x.program].college, x.program, x.year, x.block, x.track or "")):
            prog = uni.programs[sub.program]
            off = sorted({d for k in uni.classes_for(sem) if sub in k.section.subcohorts for d in k.course.placement})
            note = f"{prog.name}, Year {sub.year}" + (f" | placement days: {', '.join(off)}" if off else "")
            row = _grid(ws, row, sub.label, note, by_sub[sub], uni, "section")

        ws = wb.create_sheet(f"Sem {sem} Teachers")
        _setup_grid_sheet(ws, uni)
        by_fac = defaultdict(list)
        for s in sched.sessions:
            if s.klass.faculty:
                by_fac[s.klass.faculty.key].append(s)
        summary = faculty_load_summary(uni, sem)
        row = 1
        for key in sorted(by_fac, key=lambda k: (uni.faculty[k].college, uni.faculty[k].department, k)):
            f = uni.faculty[key]
            r = summary[key]
            note = f"{f.rank}, {f.department} ({f.college}) | {r['lu']} LU of {f.max_lu}, {r['classes']} classes"
            row = _grid(ws, row, f.name, note, by_fac[key], uni, "teacher")

        ws = wb.create_sheet(f"Sem {sem} Rooms")
        _setup_grid_sheet(ws, uni)
        by_room = defaultdict(list)
        for s in sched.sessions:
            if s.room:
                by_room[s.room].append(s)
        row = 1
        pool_of = {r.name: r.pool for r in uni.campus.rooms}
        for room in sorted(by_room, key=lambda r: (pool_of.get(r, ""), r)):
            row = _grid(ws, row, room, pool_of.get(room, ""), by_room[room], uni, "room")

        ws = wb.create_sheet(f"Sem {sem} List")
        _table(ws, ["Course", "Title", "Section", "College", "Part", "Days", "Time", "Pattern", "Room",
                    "Teacher", "Teacher's unit", "Type", "Needs approval"],
               _meeting_rows(uni, sched), [10, 40, 12, 8, 11, 8, 12, 10, 16, 32, 10, 7, 30])

    # Loads.
    ws = wb.create_sheet("Loads")
    rows = []
    for key, f in uni.faculty.items():
        row = [f.college, f.department, f.name, f.rank, f.designation, f.max_lu]
        for sem, sched in schedules.items():
            r = faculty_load_summary(uni, sem)[key]
            days = defaultdict(int)
            for s in sched.sessions:
                if s.klass.faculty is f and s.part != "NSTP":
                    days[s.day] += s.end - s.start
            busiest = max(days.values(), default=0) / 60
            row += [r["lu"], f"{r['preps_capped']}/{r['preps_total']}", r["classes"], len(days), round(busiest, 1)]
        rows.append(row)
    headers = ["College", "Department", "Name", "Rank", "Designation", "Max LU"]
    for sem in schedules:
        headers += [f"S{sem} LU", f"S{sem} Preps", f"S{sem} Classes", f"S{sem} Days", f"S{sem} Longest day (h)"]
    _table(ws, headers, rows, [8, 24, 30, 22, 30, 7] + [8, 8, 8, 7, 10] * len(schedules))

    # Room use.
    ws = wb.create_sheet("Room use")
    weekly = 39 * 60
    rows = []
    for r in uni.campus.rooms:
        row = [r.name, r.pool]
        for sem, sched in schedules.items():
            mins = sum(s.end - s.start for s in sched.sessions if s.room == r.name and s.day != "Sat")
            row.append(round(100 * mins / (weekly * r.capacity)))
        rows.append(row)
    _table(ws, ["Room", "Pool"] + [f"S{sem} % of 39 h" for sem in schedules], rows, [22, 28, 14, 14])

    # Needs approval.
    ws = wb.create_sheet("Needs approval")
    rows = []
    for sem, sched in schedules.items():
        for r in _meeting_rows(uni, sched):
            if r[-1]:
                rows.append([sem] + r[:3] + [r[5], r[6], r[8], r[9], r[-1]])
    _table(ws, ["Sem", "Course", "Title", "Section", "Days", "Time", "Room", "Teacher", "Reason"], rows,
           [5, 10, 40, 12, 8, 12, 12, 30, 50])

    ws = wb.create_sheet("Plan notes")
    rows = [[i.level, i.category, i.message] for i in validate(uni)]
    _table(ws, ["Level", "Topic", "Note"], rows, [9, 10, 150])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
